import os
import asyncio
import json
import mcp.types as types
import mcp.server.stdio
from mcp.server import Server, NotificationOptions
from mcp.server.models import InitializationOptions
from mcp.shared.exceptions import McpError
from pydantic import BaseModel, Field
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document



# helpers
def get_MIME_type(path: str) -> str:
    '''
    Get the MIME type of a file based on its extension.

    Args:
    path (str): The file path.

    Returns:
    str: The MIME type of the file.
    '''

    ext = os.path.splitext(path)[1].lower()
    mime_types = {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".py": "text/x-python",
        ".js": "text/javascript",
        ".json": "application/json",
        ".html": "text/html",
        ".css": "text/css",
        ".csv": "text/csv",
        ".xml": "application/xml",
        ".yaml": "application/x-yaml",
        ".yml": "application/x-yaml",
        '.pdf': 'application/pdf',
        '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        '.ppt': 'application/vnd.ms-powerpoint',
        '.doc': 'application/msword',
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    }
    return mime_types.get(ext, "application/octet-stream")


def is_safe_path(path: str, root_path: str) -> bool:
    """
    Check if the given path is a safe path within the root path.

    Args:
    path (str): The path to check.
    root_path (str): The root path.

    Returns:
    bool: True if the path is safe, False otherwise.
    """

    if not root_path:
        return False
    

    abs_path = os.path.abspath(os.path.join(root_path, path))
    if not abs_path.startswith(root_path):
        return False
    return True

def read_pdf(path: str) -> str:
    '''
    Read the content of a PDF file.
    
    Args:
    path (str): The path to the PDF file.

    Returns:
    str: The text content of the PDF file.
    '''
    pdf_loader = PdfReader(open(path, "rb"))
    pdf_text = ""
    for page_num in range(len(pdf_loader.pages)):
        pdf_page = pdf_loader.pages[page_num]
        pdf_text += pdf_page.extract_text()
    return pdf_text

def read_pptx(path: str) -> str:
    '''
    Read the content of a PowerPoint file.

    Args:
    path (str): The path to the PowerPoint file.

    Returns:
    str: The text content of the PowerPoint file.
    '''
    presentation = Presentation(path)
    text = ""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def read_docx(path: str) -> str:
    '''
    Read the content of a Word document.

    Args:
    path (str): The path to the Word document.

    Returns:
    str: The text content of the Word document.
    '''
    doc = Document(path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text
 

    
# Server initialization ---------------------------------
async def run_server(
        root_path: str,

) -> None:
    """
    Run the server with the given root path.
    
    Args:
    root_path (str): The root path for the server.
    """
    if not os.path.exists(root_path):
        raise ValueError(f"Root path {root_path} does not exist.")
    
    root_path = os.path.abspath(root_path)

    # server
    server = Server('filesystem')

    # resources
    @server.list_resources()
    async def list_resources() -> list[types.Resource]:
        """
        List all resources in the root path.
        
        Returns:
        list[types.Resource]: List of resources in the root path.
        """
        resources = []
        for root, _, files in os.walk(root_path):
            for file in files:
                full_path = os.path.join(root, file)
                relpath = os.path.relpath(full_path, root_path)
                uri = f'file://{relpath}'
                resources.append(types.Resource(
                    uri=uri,
                    name=relpath,
                    description=f"File: {file}", 
                    path=full_path,
                    mimeType= get_MIME_type(full_path)
                    ))
        return resources
    
    @server.read_resource()
    async def read_resource(uri: types.AnyUrl) -> str:
        """
        Read the content of a resource.
        
        Args:
        uri (types.AnyUrl): The URI of the resource.
        
        Returns:
        str: The content of the resource.

        Raises:
        McpError: If the resource does not exist.
        """

        if uri.scheme != 'file':
            raise McpError(
                types.ErrorData(
                    code= types.INVALID_PARAMS, 
                    message= f"Invalid URI scheme: {uri.scheme}. Expected 'file'."
                )
            )
        
        path = str(uri).replace('file://', '', 1)

        if not is_safe_path(path, root_path):
            raise McpError(
                types.ErrorData(
                    code= types.INVALID_PARAMS, 
                    message= f"Unsafe path: {path}. Path must be within the root path."
                )
            )
        
        full_path = os.path.join(root_path, path)
        if not os.path.exists(full_path) or not os.path.isfile(full_path):
            raise McpError(
                types.ErrorData(
                    code= types.INVALID_PARAMS, 
                    message= f"Resource does not exist: {full_path}."
                )
            )
        
        # will be handle multiple file types 
        try:
            with open(full_path, 'rb') as file:
                content = file.read()
        except Exception as e:
            raise McpError(
                types.ErrorData(
                    code= types.INTERNAL_ERROR, 
                    message= f"Failed to read resource: {full_path}. Error: {str(e)}"
                )
            )


        

