import os 
import mcp.types as types
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document



# root_path = os.path.abspath('./resourses')
# for root, _, files in os.walk(root_path):
#     print('root:', root)
#     for file in files:
#         full_path = os.path.join(root, file)
#         relpath = os.path.relpath(full_path, root_path)
#         uri = f'file://{relpath}'
#         print('relpath:', relpath)
#         print('full_path:', full_path)
#         print('ext:', os.path.splitext(relpath)[1].lower())
#         #print('schema:', types.AnyUrl(uri).scheme)

# ext = '.txt'
# mime_types = {
#         ".txt": "text/plain",
#         ".md": "text/markdown",
#         ".py": "text/x-python",
#         ".js": "text/javascript",
#         ".json": "application/json",
#         ".html": "text/html",
#         ".css": "text/css",
#         ".csv": "text/csv",
#         ".xml": "application/xml",
#         ".yaml": "application/x-yaml",
#         ".yml": "application/x-yaml",
#     }

# print(mime_types.get(ext, "application/octet-stream"))

# full_path = './resourses/Digital Multimedia/Lecture 1.pdf'
# with open(full_path, "rb") as f:
#     print(f.read())

# def read_pdf(path: str) -> str:
#     pdf_loader = PdfReader(open(path, "rb"))
#     pdf_text = ""
#     for page_num in range(len(pdf_loader.pages)):
#         pdf_page = pdf_loader.pages[page_num]
#         pdf_text += pdf_page.extract_text()
#     return pdf_text

# print(read_pdf('./resourses/Digital Multimedia/Lecture 1.pdf'))

# def read_text_based_files(path: str) -> str:
#     with open(path, "r", encoding="utf-8") as f:
#         return f.read()
    
# print(read_text_based_files('test_csv.csv'))

# read presentations
# def read_presentation(path: str) -> str:
#     presentation = Presentation(path)
#     text = ""
#     for slide in presentation.slides:
#         for shape in slide.shapes:
#             if hasattr(shape, "text"):
#                 text += shape.text + "\n"
#     return text

# print(read_presentation('pptexamples.ppt'))

# read documents
def read_document(path: str) -> str:
    doc = Document(path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

print(read_document('demo.docx'))